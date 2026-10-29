"""Export the user's serial/Phrase workbook verbatim, without evaluating cells."""
from pathlib import Path
from openpyxl import load_workbook
import hashlib
import json

ROOT=Path(__file__).resolve().parents[2]
source=ROOT/'DataTable/en-gb/Language.xlsx'
book=load_workbook(source,read_only=True,data_only=False)
sheet=book['NDLanguage']
entries={}
for row,(serial,phrase) in enumerate(sheet.iter_rows(min_row=3,values_only=False),3):
    if serial.data_type=='f' or phrase.data_type=='f':raise ValueError(f'Formula at row {row}')
    key,text=serial.value,phrase.value
    if not isinstance(key,int) or key<0 or str(key) in entries or not isinstance(text,str):
        raise ValueError(f'Invalid or duplicate language row {row}')
    entries[str(key)]=text
result=dict(source='DataTable/en-gb/Language.xlsx',sheet=sheet.title,sha256=hashlib.sha256(source.read_bytes()).hexdigest(),count=len(entries),phrases=entries)
out=ROOT/'Map/Platform01/assets/language.json'
out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(f'Exported {len(entries)} verbatim phrases; {out.stat().st_size} bytes')
